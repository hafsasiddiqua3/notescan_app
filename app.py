
import streamlit as st
import anthropic
import base64
from PIL import Image
import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import tempfile
import os

# Page config
st.set_page_config(
    page_title="NoteScan AI - OCR with Diagrams",
    page_icon="📝",
    layout="wide"
)

# Custom CSS
st.markdown("""
    <style>
    .main-header {
        font-size: 2.5rem;
        font-weight: bold;
        color: #1f77b4;
        text-align: center;
        margin-bottom: 1rem;
    }
    .sub-header {
        text-align: center;
        color: #666;
        margin-bottom: 2rem;
    }
    </style>
""", unsafe_allow_html=True)

# Header
st.markdown('<p class="main-header">📝 NoteScan AI</p>', unsafe_allow_html=True)
st.markdown('<p class="sub-header">Convert handwritten notes & diagrams to editable documents</p>', unsafe_allow_html=True)

# Sidebar
with st.sidebar:
    st.header("⚙️ Settings")
    api_key = st.text_input("Anthropic API Key", type="password", help="Get your API key from console.anthropic.com")
    
    st.markdown("---")
    st.subheader("📋 Instructions")
    st.markdown("""
    1. Enter your Anthropic API key
    2. Upload an image of handwritten notes
    3. Choose output format
    4. Click 'Process Image'
    5. Download your editable document!
    
    **Supports:**
    - ✅ Handwritten text
    - ✅ Diagrams & drawings
    - ✅ Chemical equations
    - ✅ Mathematical notation
    - ✅ Tables & lists
    """)
    
    st.markdown("---")
    st.info("💡 Tip: Clear, well-lit photos work best!")

# Main content
col1, col2 = st.columns([1, 1])

with col1:
    st.subheader("📤 Upload Image")
    uploaded_file = st.file_uploader(
        "Choose an image file",
        type=['png', 'jpg', 'jpeg', 'webp'],
        help="Upload a clear photo of your handwritten notes"
    )
    
    if uploaded_file:
        image = Image.open(uploaded_file)
        st.image(image, caption="Uploaded Image", use_container_width=True)
        
        # Output format selection
        output_format = st.radio(
            "Select output format:",
            ["Markdown (.md)", "Word Document (.docx)", "Plain Text (.txt)"],
            horizontal=True
        )

with col2:
    st.subheader("📄 Extracted Content")
    
    if uploaded_file and api_key:
        if st.button("🚀 Process Image", type="primary", use_container_width=True):
            with st.spinner("🔍 Analyzing image with AI..."):
                try:
                    # Convert image to base64
                    img_byte_arr = io.BytesIO()
                    image.save(img_byte_arr, format=image.format or 'PNG')
                    img_byte_arr = img_byte_arr.getvalue()
                    base64_image = base64.b64encode(img_byte_arr).decode('utf-8')
                    
                    # Determine media type
                    media_type_map = {
                        'PNG': 'image/png',
                        'JPEG': 'image/jpeg',
                        'JPG': 'image/jpeg',
                        'WEBP': 'image/webp'
                    }
                    media_type = media_type_map.get(image.format or 'PNG', 'image/png')
                    
                    # Call Claude API
                    client = anthropic.Anthropic(api_key=api_key)
                    
                    message = client.messages.create(
                        model="claude-sonnet-4-20250514",
                        max_tokens=4096,
                        messages=[
                            {
                                "role": "user",
                                "content": [
                                    {
                                        "type": "image",
                                        "source": {
                                            "type": "base64",
                                            "media_type": media_type,
                                            "data": base64_image,
                                        },
                                    },
                                    {
                                        "type": "text",
                                        "text": """Please analyze this handwritten note image and extract ALL content in a structured format.

Instructions:
1. Extract ALL handwritten text accurately, preserving spelling and formatting
2. Describe any diagrams, drawings, or visual elements in detail
3. Preserve the structure (headings, lists, sections)
4. Include mathematical equations and chemical formulas
5. Maintain the hierarchical organization

Format your response as:
# [Title if present]

## [Section headings]
- Transcribed text content
- Lists and bullet points

### Diagrams/Drawings:
[Detailed descriptions of any visual elements]

### Equations/Formulas:
[Chemical equations, math formulas]

Be thorough and accurate. Include everything visible in the image."""
                                    }
                                ],
                            }
                        ],
                    )
                    
                    # Extract response
                    extracted_text = message.content[0].text
                    
                    # Display in app
                    st.success("✅ Processing complete!")
                    st.markdown("### Extracted Content:")
                    st.markdown(extracted_text)
                    
                    # Prepare download based on format
                    if "Markdown" in output_format:
                        st.download_button(
                            label="📥 Download Markdown",
                            data=extracted_text,
                            file_name="extracted_notes.md",
                            mime="text/markdown",
                            use_container_width=True
                        )
                    
                    elif "Word Document" in output_format:
                        # Create Word document
                        doc = Document()
                        
                        # Add title
                        title = doc.add_heading('Extracted Notes', 0)
                        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        
                        # Process markdown-style content
                        lines = extracted_text.split('\n')
                        for line in lines:
                            line = line.strip()
                            if not line:
                                continue
                            
                            if line.startswith('# '):
                                doc.add_heading(line[2:], level=1)
                            elif line.startswith('## '):
                                doc.add_heading(line[3:], level=2)
                            elif line.startswith('### '):
                                doc.add_heading(line[4:], level=3)
                            elif line.startswith('- ') or line.startswith('* '):
                                p = doc.add_paragraph(line[2:], style='List Bullet')
                            else:
                                doc.add_paragraph(line)
                        
                        # Save to bytes
                        doc_bytes = io.BytesIO()
                        doc.save(doc_bytes)
                        doc_bytes.seek(0)
                        
                        st.download_button(
                            label="📥 Download Word Document",
                            data=doc_bytes,
                            file_name="extracted_notes.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            use_container_width=True
                        )
                    
                    else:  # Plain text
                        # Remove markdown formatting for plain text
                        plain_text = extracted_text.replace('#', '').replace('*', '').replace('_', '')
                        st.download_button(
                            label="📥 Download Plain Text",
                            data=plain_text,
                            file_name="extracted_notes.txt",
                            mime="text/plain",
                            use_container_width=True
                        )
                    
                except anthropic.AuthenticationError:
                    st.error("❌ Invalid API key. Please check your Anthropic API key.")
                except Exception as e:
                    st.error(f"❌ Error: {str(e)}")
    
    elif uploaded_file and not api_key:
        st.warning("⚠️ Please enter your Anthropic API key in the sidebar to process the image.")
    
    else:
        st.info("👆 Upload an image to get started!")

# Footer
st.markdown("---")
st.markdown("""
<div style='text-align: center; color: #666; padding: 20px;'>
    <p>Built with ❤️ using Streamlit & Claude AI</p>
    <p style='font-size: 0.8rem;'>Supports handwriting, diagrams, equations, and more!</p>
</div>
""", unsafe_allow_html=True)